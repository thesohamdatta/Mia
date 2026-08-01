import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import re
import os

# Create workbook
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "YouTube Transcripts Catalog"

# Define styles
header_font = Font(bold=True, size=11, color="FFFFFF")
header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)
wrap_alignment = Alignment(wrap_text=True, vertical='top')

# Headers
headers = [
    "ID", "YouTube ID", "Title", "Speaker(s)", "Channel", 
    "Views", "Category", "Subcategory", "Key Topics", 
    "Context/Relevance to MIA", "Transcript File", "YouTube URL", "Status"
]

for col, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = thin_border

# Set column widths
col_widths = [5, 12, 50, 30, 18, 10, 25, 25, 40, 40, 55, 35, 10]
for i, width in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = width

# Parse transcript files from the directory structure
transcript_root = r"C:\Users\Soham\Downloads\AI\archive\TRANSCRIPT"
all_transcripts = []

# Category mapping based on directory structure
category_map = {
    "ai-engineer-archive": {
        "transcripts": "General AI Engineering Talks",
        "agent-engineering-and-evals": "Agent Engineering & Evaluation",
        "anthropic-and-fde": "Anthropic & Forward Deployed Engineering",
        "coding-tools-and-workflows": "Coding Tools & Workflows",
        "context-layer-and-long-horizon": "Context Layer & Long-Horizon",
        "developer-productivity-and-architecture": "Developer Productivity & Architecture",
        "harnesses-and-scaffolding": "Harnesses & Scaffolding",
        "openai-and-keynotes": "OpenAI & Keynotes",
        "reliable-agents-and-patterns": "Reliable Agents & Patterns",
        "security-auth-and-tool-calling": "Security, Auth & Tool Calling",
        "advanced-agent-design": "Advanced Agent Design"
    },
    "AIengg": {
        "agent-engineering-and-evals": "Agent Engineering & Evaluation",
        "anthropic-and-fde": "Anthropic & Forward Deployed Engineering",
        "coding-tools-and-workflows": "Coding Tools & Workflows",
        "context-layer-and-long-horizon": "Context Layer & Long-Horizon",
        "developer-productivity-and-architecture": "Developer Productivity & Architecture",
        "harnesses-and-scaffolding": "Harnesses & Scaffolding",
        "openai-and-keynotes": "OpenAI & Keynotes",
        "reliable-agents-and-patterns": "Reliable Agents & Patterns",
        "security-auth-and-tool-calling": "Security, Auth & Tool Calling",
        "advanced-agent-design": "Advanced Agent Design"
    }
}

def extract_youtube_id(filename):
    """Extract YouTube ID from filename"""
    match = re.search(r'(?:^\d+_)?([a-zA-Z0-9_-]{11})_', filename)
    if match:
        return match.group(1)
    match = re.search(r'^([a-zA-Z0-9_-]{11})_', filename)
    if match:
        return match.group(1)
    return ""

def extract_title(filename):
    """Extract title from filename"""
    title = re.sub(r'^\d+_', '', filename)
    title = re.sub(r'^[a-zA-Z0-9_-]{11}_', '', title)
    title = title.replace('.md', '')
    title = title.replace('_', ' ')
    title = re.sub(r'\s+', ' ', title).strip()
    return title

def extract_speaker_from_content(filepath):
    """Try to extract speaker from transcript content"""
    try:
        with open(filepath, 'r', encoding='utf-8') as f:
            content = f.read(2000)
        speaker_match = re.search(r'(?:Speaker|Presenter):\s*([^\n\r]+)', content, re.IGNORECASE)
        if speaker_match:
            return speaker_match.group(1).strip()
        speaker_match = re.search(r'\*\*Speaker\*\*:\s*([^\n\r]+)', content)
        if speaker_match:
            return speaker_match.group(1).strip()
        speaker_match = re.search(r'-\s*\*\*Speaker\*\*:\s*([^\n\r]+)', content)
        if speaker_match:
            return speaker_match.group(1).strip()
    except:
        pass
    return ""

def extract_views_from_content(filepath):
    """Try to extract views from transcript content"""
    try:
        with open(filepath, 'r', encoding='utf-8') as f:
            content = f.read(2000)
        views_match = re.search(r'(?:Views|Popularity).*?(\d+(?:,\d+)*(?:k|K)?)', content, re.IGNORECASE)
        if views_match:
            return views_match.group(1)
    except:
        pass
    return ""

def get_category_and_subcategory(root_path, filepath):
    """Determine category and subcategory from file path"""
    rel_path = os.path.relpath(filepath, root_path)
    parts = rel_path.split(os.sep)
    
    if len(parts) >= 3:
        main_cat = parts[0]
        sub_cat = parts[1]
        
        # Handle ai-engineer-archive-v2 (duplicate of ai-engineer-archive)
        if main_cat == "ai-engineer-archive-v2":
            if main_cat == "ai-engineer-archive-v2" and sub_cat == "transcripts":
                category = "General AI Engineering Talks"
                return category, "Transcripts"
            elif sub_cat in category_map.get("ai-engineer-archive", {}):
                cat_info = category_map["ai-engineer-archive"]
                category = cat_info.get(sub_cat, "General")
                return category, sub_cat.replace('-', ' ').title()
            return "General", "General"
        
        # Handle ai-engineer-archive
        if main_cat == "ai-engineer-archive" and sub_cat == "transcripts":
            category = "General AI Engineering Talks"
            return category, "Transcripts"
        elif main_cat == "AIengg":
            if sub_cat in category_map.get("AIengg", {}):
                cat_info = category_map["AIengg"]
                category = cat_info.get(sub_cat, "General")
                return category, sub_cat.replace('-', ' ').title()
            return "General", "General"
        
        if main_cat in category_map:
            cat_info = category_map[main_cat]
            category = cat_info.get(sub_cat, cat_info.get("transcripts", "General"))
            return category, sub_cat.replace('-', ' ').title()
    
    return "General", "General"

def get_mia_relevance(category, subcategory, title):
    """Determine relevance to MIA based on category and title"""
    relevance_map = {
        "Agent Engineering": "Direct reference for MIA's agent/skill architecture",
        "Agent Evaluation": "Critical for MIA's eval framework and constitution-check",
        "Verification": "MIA's health gate and adversarial review based on this",
        "Context Engineering": "MIA's learning layer and daily rituals",
        "Agent Architecture": "MIA's daemon + skill system design",
        "Multi-Agent": "MIA's AURA: single agent + dynamic sub-agents",
        "Loops/Automation": "MIA's grill-to-ship = disciplined loops",
        "Reliable Agents": "MIA's daemon resilience and health checks",
        "Coding Workflows": "MIA's grill-to-ship pipeline reference",
        "Agent Memory": "MIA's learning layer: JSONL, decay, memory.md",
        "Local-First AI": "MIA's core principle: local-first, compiled binaries",
        "Software Factories": "MIA's philosophy: constitutional runtime > harness only",
        "Long-Horizon": "MIA's daemon architecture: brain/hands decoupling",
        "Anthropic": "MIA's constitutional principles from Anthropic",
    }
    
    for key, relevance in relevance_map.items():
        if key.lower() in category.lower() or key.lower() in subcategory.lower() or key.lower() in title.lower():
            return relevance
    
    return f"General reference for MIA's {category.lower()} knowledge base"

def process_transcript_files():
    """Walk through all transcript files and extract metadata"""
    transcript_files = []
    
    for root, dirs, files in os.walk(transcript_root):
        for file in files:
            if file.endswith('.md') and not file.startswith('README') and not file.startswith('INDEX'):
                full_path = os.path.join(root, file)
                transcript_files.append(full_path)
    
    print(f"Found {len(transcript_files)} transcript files")
    
    for idx, filepath in enumerate(transcript_files, 1):
        filename = os.path.basename(filepath)
        youtube_id = extract_youtube_id(filename)
        title = extract_title(filename)
        
        if not youtube_id:
            try:
                with open(filepath, 'r', encoding='utf-8') as f:
                    content = f.read(1000)
                yt_match = re.search(r'(?:youtube\.com/watch\?v=|youtu\.be/)([a-zA-Z0-9_-]{11})', content)
                if yt_match:
                    youtube_id = yt_match.group(1)
            except:
                pass
        
        speaker = extract_speaker_from_content(filepath)
        views = extract_views_from_content(filepath)
        category, subcategory = get_category_and_subcategory(transcript_root, filepath)
        mia_relevance = get_mia_relevance(category, subcategory, title)
        rel_path = os.path.relpath(filepath, r"C:\Users\Soham\Downloads\AI")
        youtube_url = f"https://www.youtube.com/watch?v={youtube_id}" if youtube_id else ""
        
        all_transcripts.append({
            'id': idx,
            'youtube_id': youtube_id,
            'title': title,
            'speaker': speaker,
            'channel': "AI Engineer" if "ai-engineer" in filepath.lower() else "AI Engineer",
            'views': views,
            'category': category,
            'subcategory': subcategory,
            'key_topics': f"{category} - {subcategory}",
            'mia_relevance': mia_relevance,
            'transcript_file': rel_path,
            'youtube_url': youtube_url,
            'status': "Archived"
        })
    
    return all_transcripts

# Process all files
print("Processing transcript files...")
all_transcripts = process_transcript_files()

# Write to Excel
for row_idx, t in enumerate(all_transcripts, 2):
    row_data = [
        t['id'], t['youtube_id'], t['title'], t['speaker'], t['channel'],
        t['views'], t['category'], t['subcategory'], t['key_topics'],
        t['mia_relevance'], t['transcript_file'], t['youtube_url'], t['status']
    ]
    for col_idx, value in enumerate(row_data, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = wrap_alignment
        cell.border = thin_border
        if row_idx % 2 == 0:
            cell.fill = PatternFill(start_color="F2F7FC", end_color="F2F7FC", fill_type="solid")

# Freeze panes and auto-filter
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(all_transcripts)+1}"

# Save
output_path = r"C:\Users\Soham\Downloads\AI\mia\YOUTUBE_TRANSCRIPTS_CATALOG_FULL.xlsx"
wb.save(output_path)
print(f"Saved to {output_path}")
print(f"Total transcripts cataloged: {len(all_transcripts)}")